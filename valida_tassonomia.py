"""
VALIDA TASSONOMIA - SO FOOD
============================================================
Incrocia il database vettoriale ChromaDB (le schede prodotto scritte a mano,
progetto "chatbot" di Nino) con il file Excel "tassonomia-prodotti-da-validare.xlsx"
per proporre automaticamente la SOTTOCATEGORIA nella colonna "-> CORREGGI QUI",
per i soli prodotti di cui esiste gia' una scheda nel DB vettoriale.

COME FUNZIONA
- Per ogni riga dell'Excel, cerca il CODICE nel DB vettoriale (match esatto,
  eventualmente disambiguato con il nome FORNITORE se il codice e' ambiguo).
- Se trova una scheda, manda il testo al modello Gemini insieme all'ELENCO
  CHIUSO delle sottocategorie ammesse (foglio "Elenco sottocategorie") e gli
  chiede di scegliere ESATTAMENTE una di quelle stringhe (mai testo libero,
  per rispettare la validazione a tendina della colonna G).
- Scrive il risultato SOLO nella colonna "-> CORREGGI QUI" (G), lasciando
  intatte REPARTO / CATEGORIA / SOTTOCATEGORIA PROPOSTA (sono la proposta da
  verificare, non va toccata) e aggiunge un commento alla cella con la
  motivazione, cosi' resta tracciabile.
- E' idempotente: se rilanciato, salta le righe gia' compilate nell'output.

COSA DEVI CONFIGURARE QUI SOTTO
- PERCORSO_EXCEL_INPUT: il file .xlsx da validare
- PERCORSO_EXCEL_OUTPUT: dove salvare il risultato (di default un file nuovo,
  non sovrascrive l'originale)
- PERCORSO_DATABASE_VETTORIALE: la cartella ./database_vettoriale del progetto chatbot
- CONFIDENZE_DA_PROCESSARE: quali righe processare (di default quelle che
  secondo il foglio stesso servono di piu': "DA ASSEGNARE" e "non classificato")

REQUISITI
    pip install chromadb google-genai pandas openpyxl python-dotenv
Un file .env nella stessa cartella con:
    GEMINI_API_KEY=la_tua_chiave
"""

import os
import sys
import time
import logging
from pathlib import Path

import pandas as pd
import chromadb
from google import genai
from google.genai import types
from dotenv import load_dotenv
import openpyxl

# Silenzia il warning innocuo "Direct use of automatic function calling (AFC)...":
# non usiamo function calling in questo script, e' solo un avviso informativo dell'SDK.
logging.getLogger("google_genai.models").setLevel(logging.ERROR)

# ====================================================================
# CONFIGURAZIONE - modifica questi valori
# ====================================================================
PERCORSO_EXCEL_INPUT = "tassonomia-prodotti-da-validare.xlsx"
PERCORSO_EXCEL_OUTPUT = "tassonomia-prodotti-validata.xlsx"
PERCORSO_DATABASE_VETTORIALE = "./database_vettoriale"
NOME_COLLEZIONE = "catalogo_sofood"
MODELLO_GEMINI = "models/gemini-3.5-flash-lite"

# Quali righe processare: solo quelle dove l'aiuto serve davvero, secondo le
# istruzioni del foglio stesso. Aggiungi "media" o "alta" se vuoi che il
# modello ricontrolli anche le proposte gia' fatte (in quel caso scrivera'
# nella colonna G solo se propone qualcosa di DIVERSO dalla proposta attuale).
CONFIDENZE_DA_PROCESSARE = {"DA ASSEGNARE", "non classificato"}

PAUSA_TRA_CHIAMATE_SEC = 5.0  # con un limite di 15 RPM il minimo teorico e' 60/15=4s: 5s da' un margine di sicurezza
SALVA_OGNI_N_PRODOTTI = 5     # salva il file su disco ogni N proposte scritte
MAX_TENTATIVI_RATE_LIMIT = 5  # se arriva comunque un errore di rate limit, riprova con attesa crescente

# ====================================================================
load_dotenv()
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
if not GEMINI_API_KEY:
    raise ValueError("GEMINI_API_KEY non trovata. Crea un file .env con GEMINI_API_KEY=... "
                      "nella stessa cartella di questo script.")


def normalizza(testo) -> str:
    return str(testo).strip().upper() if testo is not None else ""


def varianti_codice(codice: str) -> set:
    """Genera le varianti di un codice prodotto per gestire i caratteri che
    caricaprodotti.py converte per compatibilita' con i nomi di cartella
    (es. '/' non ammesso su Windows, sostituito con '__').
    Stessa logica di _varianti_normalizzazione_codice in caricaprodotti.py,
    cosi' un codice come 'CAPREA1/2' in Excel matcha la cartella 'CAPREA1__2'."""
    codice = normalizza(codice)
    varianti = {codice}
    varianti.add(codice.replace("_", "."))
    varianti.add(codice.replace(".", "_"))
    varianti.add(codice.replace("__", "/"))
    varianti.add(codice.replace("/", "__"))
    return varianti


def carica_vocabolario(percorso_excel: str):
    """Legge il foglio 'Elenco sottocategorie' e costruisce il vocabolario chiuso."""
    tabella = pd.read_excel(percorso_excel, sheet_name="Elenco sottocategorie")
    tabella.columns = [str(c).strip() for c in tabella.columns]
    vocabolario = []
    sottocategorie_valide = set()
    for _, riga in tabella.iterrows():
        sott = str(riga["SOTTOCATEGORIA"]).strip()
        cat = str(riga["CATEGORIA"]).strip()
        rep = str(riga["REPARTO"]).strip()
        if sott and sott.lower() != "nan":
            vocabolario.append({"sottocategoria": sott, "categoria": cat, "reparto": rep})
            sottocategorie_valide.add(sott)
    return vocabolario, sottocategorie_valide


def carica_schede_prodotto(percorso_db: str, nome_collezione: str):
    """Legge tutte le schede dal DB vettoriale e le indicizza per codice_prodotto normalizzato
    (e per tutte le sue varianti, cosi' un codice con '/' matcha anche la versione con '__')."""
    client = chromadb.PersistentClient(path=percorso_db)
    try:
        collezione = client.get_collection(name=nome_collezione)
    except Exception as errore:
        raise RuntimeError(
            f"Impossibile aprire la collezione '{nome_collezione}' in {percorso_db}. "
            f"Hai gia' lanciato caricaprodotti.py? Dettaglio: {errore}"
        )

    dati = collezione.get()  # nessun embedding necessario per una get() semplice
    indice = {}
    for doc_id, documento, metadati in zip(dati["ids"], dati["documents"], dati["metadatas"]):
        codice_grezzo = metadati.get("codice_prodotto")
        if not codice_grezzo:
            continue
        scheda = {
            "id": doc_id,
            "documento": documento,
            "nome_fornitore": metadati.get("nome_fornitore", ""),
            "categoria_prodotto": metadati.get("categoria_prodotto", ""),
        }
        for variante in varianti_codice(codice_grezzo):
            lista = indice.setdefault(variante, [])
            if scheda not in lista:  # evita doppioni quando due varianti coincidono
                lista.append(scheda)
    codici_distinti_reali = len({normalizza(m.get("codice_prodotto")) for m in dati["metadatas"] if m.get("codice_prodotto")})
    print(f"[INFO] Schede lette dal DB vettoriale: {len(dati['ids'])} "
          f"({codici_distinti_reali} codici distinti).")
    return indice


def trova_scheda(codice_excel: str, fornitore_excel: str, indice_db: dict):
    """Cerca la scheda corrispondente al CODICE della riga Excel (provando anche le
    varianti con '/' <-> '__'), disambiguando per fornitore se serve."""
    candidati = []
    for variante in varianti_codice(codice_excel):
        for candidato in indice_db.get(variante, []):
            if candidato not in candidati:
                candidati.append(candidato)

    if not candidati:
        return None
    if len(candidati) == 1:
        return candidati[0]
    # piu' fornitori hanno lo stesso codice prodotto: prova a disambiguare per nome fornitore
    fornitore_norm = normalizza(fornitore_excel)
    for candidato in candidati:
        if normalizza(candidato["nome_fornitore"]) in fornitore_norm or fornitore_norm in normalizza(candidato["nome_fornitore"]):
            return candidato
    return None  # ambiguo, meglio saltare che sbagliare


def costruisci_prompt(descrizione_riga: str, testo_scheda: str, vocabolario: list) -> str:
    elenco_vocabolario = "\n".join(
        f"- {v['sottocategoria']}  (Categoria: {v['categoria']} | Reparto: {v['reparto']})"
        for v in vocabolario
    )
    return f"""Sei un classificatore di prodotti alimentari per un grossista (SO FOOD).
Devi assegnare la SOTTOCATEGORIA corretta a un prodotto, scegliendo UNA SOLA voce
dall'elenco chiuso qui sotto. Non inventare mai una sottocategoria fuori da questo elenco:
devi restituire il testo esattamente identico a una delle voci elencate.

ELENCO CHIUSO DELLE SOTTOCATEGORIE AMMESSE:
{elenco_vocabolario}

PRODOTTO DA CLASSIFICARE:
Descrizione a gestionale: {descrizione_riga}

Scheda prodotto (fonte primaria, fidati di questa se in conflitto con la descrizione):
{testo_scheda}

Rispondi SOLO in JSON con questo schema, senza altro testo:
{{"sottocategoria": "<una voce ESATTA dall'elenco>", "motivazione": "<una frase breve>"}}
"""


def e_errore_rate_limit(errore: Exception) -> bool:
    testo_errore = str(errore).upper()
    return "429" in testo_errore or "RESOURCE_EXHAUSTED" in testo_errore or "RATE LIMIT" in testo_errore


def classifica_con_gemini(client_genai, descrizione_riga, testo_scheda, vocabolario, sottocategorie_valide):
    prompt = costruisci_prompt(descrizione_riga, testo_scheda, vocabolario)

    for tentativo in range(1, MAX_TENTATIVI_RATE_LIMIT + 1):
        try:
            risposta = client_genai.models.generate_content(
                model=MODELLO_GEMINI,
                contents=prompt,
                config=types.GenerateContentConfig(
                    temperature=0.1,
                    response_mime_type="application/json",
                ),
            )
            break
        except Exception as errore:
            if e_errore_rate_limit(errore) and tentativo < MAX_TENTATIVI_RATE_LIMIT:
                attesa = PAUSA_TRA_CHIAMATE_SEC * (2 ** tentativo)  # backoff esponenziale
                print(f"   [RATE LIMIT] Attendo {attesa:.0f}s e riprovo (tentativo {tentativo}/{MAX_TENTATIVI_RATE_LIMIT})...")
                time.sleep(attesa)
                continue
            raise

    import json
    dati = json.loads(risposta.text)
    sottocategoria = str(dati.get("sottocategoria", "")).strip()
    motivazione = str(dati.get("motivazione", "")).strip()
    if sottocategoria not in sottocategorie_valide:
        raise ValueError(f"Sottocategoria fuori vocabolario: '{sottocategoria}'")
    return sottocategoria, motivazione


def main():
    if not Path(PERCORSO_EXCEL_OUTPUT).exists():
        percorso_di_lavoro = PERCORSO_EXCEL_INPUT
        print(f"[INFO] Nessun output precedente trovato, parto da: {percorso_di_lavoro}")
    else:
        percorso_di_lavoro = PERCORSO_EXCEL_OUTPUT
        print(f"[INFO] Trovato output precedente, riprendo da: {percorso_di_lavoro}")

    vocabolario, sottocategorie_valide = carica_vocabolario(percorso_di_lavoro)
    indice_db = carica_schede_prodotto(PERCORSO_DATABASE_VETTORIALE, NOME_COLLEZIONE)

    wb = openpyxl.load_workbook(percorso_di_lavoro)
    ws = wb["Prodotti da validare"]

    intestazioni = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    col_codice = intestazioni["CODICE"]
    col_descrizione = intestazioni["DESCRIZIONE"]
    col_fornitore = intestazioni["FORNITORE"]
    col_sott_proposta = intestazioni["SOTTOCATEGORIA PROPOSTA"]
    col_correggi = intestazioni["\u2192 CORREGGI QUI"]
    col_confidenza = intestazioni["CONFIDENZA"]

    client_genai = genai.Client(api_key=GEMINI_API_KEY)

    n_totale_righe = ws.max_row - 1
    n_matchate = 0
    n_gia_compilate = 0
    n_proposte_scritte = 0
    n_nessun_match = 0
    n_ambigue = 0
    n_errori = 0

    for riga in range(2, ws.max_row + 1):
        confidenza = ws.cell(row=riga, column=col_confidenza).value
        if confidenza not in CONFIDENZE_DA_PROCESSARE:
            continue

        cella_correggi = ws.cell(row=riga, column=col_correggi)
        if cella_correggi.value not in (None, ""):
            n_gia_compilate += 1
            continue

        codice = ws.cell(row=riga, column=col_codice).value
        descrizione = ws.cell(row=riga, column=col_descrizione).value or ""
        fornitore = ws.cell(row=riga, column=col_fornitore).value or ""

        scheda = trova_scheda(codice, fornitore, indice_db)
        if scheda is None:
            candidati_multipli = []
            for variante in varianti_codice(codice):
                for candidato in indice_db.get(variante, []):
                    if candidato not in candidati_multipli:
                        candidati_multipli.append(candidato)
            if len(candidati_multipli) > 1:
                n_ambigue += 1
            else:
                n_nessun_match += 1
            continue

        n_matchate += 1
        try:
            sottocategoria, motivazione = classifica_con_gemini(
                client_genai, descrizione, scheda["documento"], vocabolario, sottocategorie_valide
            )
        except Exception as errore:
            print(f"   [ERRORE] Riga {riga} ({codice}): {errore}")
            n_errori += 1
            time.sleep(PAUSA_TRA_CHIAMATE_SEC)
            continue

        proposta_attuale = ws.cell(row=riga, column=col_sott_proposta).value
        if sottocategoria == proposta_attuale:
            # il modello conferma la proposta esistente: non serve correggere nulla
            time.sleep(PAUSA_TRA_CHIAMATE_SEC)
            continue

        cella_correggi.value = sottocategoria
        cella_correggi.comment = openpyxl.comments.Comment(
            f"Suggerito dal DB vettoriale (scheda prodotto) + Gemini.\nMotivazione: {motivazione}",
            "valida_tassonomia.py",
        )
        n_proposte_scritte += 1
        print(f"   [OK] Riga {riga} ({codice}): -> {sottocategoria}")

        # salvataggio incrementale, cosi' non si perde lavoro se lo script
        # si interrompe (o lo fermi tu) a meta'
        if n_proposte_scritte % SALVA_OGNI_N_PRODOTTI == 0:
            wb.save(PERCORSO_EXCEL_OUTPUT)
            print(f"   [SALVATAGGIO] {n_proposte_scritte} proposte salvate su disco.")

        time.sleep(PAUSA_TRA_CHIAMATE_SEC)

    wb.save(PERCORSO_EXCEL_OUTPUT)

    print("\n" + "=" * 60)
    print("RIEPILOGO")
    print("=" * 60)
    print(f"Righe totali nel foglio:                 {n_totale_righe}")
    print(f"Righe nel range di confidenza scelto:     {n_matchate + n_nessun_match + n_ambigue + n_gia_compilate}")
    print(f"  - gia' compilate in un run precedente:  {n_gia_compilate}")
    print(f"  - matchate nel DB vettoriale:           {n_matchate}")
    print(f"      -> proposte scritte in colonna G:   {n_proposte_scritte}")
    print(f"      -> confermate (nessuna modifica):   {n_matchate - n_proposte_scritte - n_errori}")
    print(f"      -> errori di classificazione:       {n_errori}")
    print(f"  - nessuna scheda nel DB (fuori target): {n_nessun_match}")
    print(f"  - codice ambiguo (piu' fornitori):      {n_ambigue}")
    print(f"\nFile salvato in: {PERCORSO_EXCEL_OUTPUT}")


if __name__ == "__main__":
    main()